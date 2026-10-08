# ksecure/scanner.py
"""Motor de escaneo K-Secure."""

import csv
from pathlib import Path

# Librerías para extraer texto de múltiples formatos
import docx
import pandas as pd
import pdfplumber
from pptx import Presentation

from ksecure.models import Finding
from ksecure.validators.base import BaseValidator
from ksecure.validators.rut import RutValidator
from ksecure.validators.phone import PhoneValidator
from ksecure.validators.credit_card import CreditCardValidator
from ksecure.validators.email import EmailValidator

class KSecureScanner:
    """Orquesta los validadores sobre una o más fuentes de datos."""

    def __init__(self, validators: list[BaseValidator] | None = None):
        # Motor inicializado con todos los validadores activos
        self.validators = validators or [
            RutValidator(), 
            PhoneValidator(),
            CreditCardValidator(),
            EmailValidator()
        ]

    def scan_file(self, filepath: str | Path) -> list[Finding]:
        """Enrutador universal: detecta la extensión y aplica el escáner correcto."""
        path = Path(filepath)
        extension = path.suffix.lower()

        if extension == ".csv":
            return self.scan_csv(path)
        elif extension == ".docx":
            return self.scan_docx(path)
        elif extension == ".xlsx":
            return self.scan_xlsx(path)
        elif extension == ".pptx":
            return self.scan_pptx(path)
        elif extension == ".pdf":
            return self.scan_pdf(path)
        elif extension == ".txt":
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return self._scan_text(f.read(), source=path.name)
        else:
            print(f"Formato no soportado aún: {extension}")
            return []

    def scan_csv(self, path: Path) -> list[Finding]:
        """Escanea un CSV simulando una tabla de base de datos."""
        findings: list[Finding] = []
        with open(path, encoding="utf-8", errors="ignore") as fh:
            reader = csv.DictReader(fh)
            for row_idx, row in enumerate(reader, start=2):
                for column, value in row.items():
                    if not value:
                        continue
                    source = f"{path.name}:{column}:fila {row_idx}"
                    findings.extend(self._scan_text(value, source))
        return findings

    def scan_docx(self, path: Path) -> list[Finding]:
        """Escanea documentos de Word."""
        findings: list[Finding] = []
        doc = docx.Document(str(path))  # <-- str() añadido
        for i, paragraph in enumerate(doc.paragraphs):
            texto = paragraph.text.strip()
            if texto:
                source = f"{path.name}:párrafo {i+1}"
                findings.extend(self._scan_text(texto, source))
        return findings

    def scan_xlsx(self, path: Path) -> list[Finding]:
        """Escanea hojas de cálculo de Excel."""
        findings: list[Finding] = []
        excel_data = pd.read_excel(str(path), sheet_name=None, dtype=str) # <-- str() añadido
        for sheet_name, df in excel_data.items():
            # Usamos enumerate(start=2) para llevar un conteo seguro de la fila real
            for i, (row_idx, row) in enumerate(df.iterrows(), start=2):
                for col_name, value in row.items():
                    if pd.notna(value) and str(value).strip() and str(value) != "nan":
                        source = f"{path.name}:hoja '{sheet_name}':col '{col_name}':fila {i}"
                        findings.extend(self._scan_text(str(value), source))
        return findings

    def scan_pptx(self, path: Path) -> list[Finding]:
        """Escanea presentaciones de PowerPoint."""
        findings: list[Finding] = []
        prs = Presentation(str(path))  # <-- str() añadido
        for i, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                # Extraemos el texto de forma segura para calmar al linter
                texto_forma = getattr(shape, "text", "").strip()
                if texto_forma:
                    source = f"{path.name}:diapositiva {i+1}"
                    findings.extend(self._scan_text(texto_forma, source))
        return findings

    def scan_pdf(self, path: Path) -> list[Finding]:
        """Escanea archivos PDF."""
        findings: list[Finding] = []
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages):
                texto = page.extract_text()
                if texto:
                    source = f"{path.name}:página {i+1}"
                    findings.extend(self._scan_text(texto, source))
        return findings

    def _scan_text(self, content: str, source: str) -> list[Finding]:
        """Pasa el texto extraído por todos los validadores activos."""
        findings: list[Finding] = []
        for validator in self.validators:
            findings.extend(validator.find(content, source))
        return findings